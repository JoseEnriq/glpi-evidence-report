"""End-to-end tests against the simulated GLPI (simulator/fake_glpi.py): the
real code of glpi_api, glpi_files and docx_builder, with no network."""

import os

import pytest
from docx import Document

import glpi_api
from generate_report import run
from glpi_api import GlpiAccessError, open_source
from models import ORIGIN_FOLLOWUP, ORIGIN_SOLUTION, ORIGIN_TICKET
from simulator import DEMO_INSTANCE, DEMO_PERIOD
from simulator.fake_glpi import DEMO_TOKEN, PAGE_LIMIT, fake_glpi_session


@pytest.fixture(autouse=True)
def environment(monkeypatch):
    monkeypatch.setenv("GLPI_DEMO_API_USER_TOKEN", DEMO_TOKEN)
    monkeypatch.setattr(glpi_api.time, "sleep", lambda _s: None)  # no waiting on the 503 retry


def _docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def test_login_and_profile_switch():
    session, _ = fake_glpi_session(DEMO_INSTANCE.url)
    source = open_source(DEMO_INSTANCE, session)
    try:
        client = source.client
        r = client.http.get(client.api + "/getActiveProfile", headers=client.headers())
        assert r.json()["active_profile"]["name"] == "Observer"  # left Self-Service
    finally:
        source.close()


def test_profile_switch_prefers_least_privileged():
    session, adapter = fake_glpi_session(DEMO_INSTANCE.url)
    adapter.data["profiles"].insert(1, {"id": 4, "name": "Super-Admin"})
    source = open_source(DEMO_INSTANCE, session)
    try:
        client = source.client
        r = client.http.get(client.api + "/getActiveProfile", headers=client.headers())
        assert r.json()["active_profile"]["name"] == "Observer"  # not Super-Admin
    finally:
        source.close()


def test_wrong_token_is_rejected(monkeypatch):
    monkeypatch.setenv("GLPI_DEMO_API_USER_TOKEN", "wrong-token")
    session, _ = fake_glpi_session(DEMO_INSTANCE.url)
    with pytest.raises(GlpiAccessError, match="rejected the login"):
        open_source(DEMO_INSTANCE, session)


def test_search_follows_pagination_and_skips_trash_and_other_month():
    session, _ = fake_glpi_session(DEMO_INSTANCE.url)
    source = open_source(DEMO_INSTANCE, session)
    try:
        ids = source.find_ticket_ids(*DEMO_PERIOD)
    finally:
        source.close()
    assert len(ids) > PAGE_LIMIT  # only passes if the client follows Content-Range
    assert ids == list(range(101, 114))  # no 99 (July) and no 114 (trash)


def test_attachments_come_from_all_three_origins():
    session, _ = fake_glpi_session(DEMO_INSTANCE.url)
    source = open_source(DEMO_INSTANCE, session)
    try:
        (t101,) = source.fetch_tickets([101])
        (t102,) = source.fetch_tickets([102])
    finally:
        source.close()
    assert {a.origin for a in t101.attachments} == {ORIGIN_TICKET, ORIGIN_SOLUTION}
    assert {a.origin for a in t102.attachments} == {ORIGIN_FOLLOWUP, ORIGIN_SOLUTION}
    assert t101.status == "Solucionado" and t101.category == "Acessos > Desbloqueio de usuário"


def test_full_generation_and_cache(tmp_path):
    cache = tmp_path / "cache"
    output = tmp_path / "report.docx"
    for _ in range(2):  # the second run must use the cache
        session, _server = fake_glpi_session(DEMO_INSTANCE.url)
        run(DEMO_INSTANCE, period=DEMO_PERIOD, technicians=DEMO_INSTANCE.technicians, output=str(output),
            http_session=session, cache_dir=str(cache))

    text = _docx_text(output)
    assert "#101" in text and "#111" in text
    assert "#112" not in text  # handled by a technician outside the list
    assert "#113" not in text  # no technician
    assert "Relatório de vendas" in text  # #105 made it despite the transient 503
    assert 'Anexo "print_tela_usuario.png" sem imagem disponível' in text  # HTML rejected as an image
    files = os.listdir(cache / "demo")
    assert len(files) == 12 and not any(f.startswith("514_") for f in files)
