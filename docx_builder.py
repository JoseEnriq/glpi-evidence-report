"""
docx_builder.py
Builds the evidence report as a Word document (.docx) from the data already
extracted from GLPI (see models.py — Ticket). This module does not talk to
GLPI: it receives the list of tickets and the map of already downloaded
images, and only lays out the document.

Layout: A4, 2 cm margins, one accent color, a cover page with a summary per
technician, a table of contents with internal links, each ticket with a
metadata grid and a colored status badge, updates as a timeline (the
solution highlighted in green), framed images with captions below, and a
header/footer with "Página X de Y".

The report itself is in Portuguese (it is the product, written for a
Brazilian audience); the code is in English.

How it is built: python-docx covers the basics (paragraphs, tables,
styles, pictures). What it does not expose — cell shading, paragraph
borders, PAGE/NUMPAGES fields, bookmarks and internal hyperlinks — is done
by creating the Word XML elements directly (OOXML, `w:` prefix) with
OxmlElement. In those cases the ORDER of child elements matters (Word
rejects the file if it does not follow the schema order), hence
insert_element_before.
"""

import calendar
from datetime import date, datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

from models import GLPI_STATUS, STATUS_CLOSED, STATUS_SOLVED

# ---------------------------------------------------------------------------
# Sizes and colors
# ---------------------------------------------------------------------------

USABLE_WIDTH_CM = 17  # A4 (21 cm) minus 2 cm margins on each side
MAX_IMAGE_WIDTH_CM = 17
MAX_IMAGE_HEIGHT_CM = 19  # keeps a very tall/narrow screenshot from taking several pages on its own

FONT = "Calibri"
ACCENT = "1F3864"  # dark navy: titles, rules, table headers
MUTED_TEXT = "595959"
LABEL_FILL = "F2F2F2"
SOFT_BORDER = "D9D9D9"
SOLUTION_GREEN = "2E7D32"
SOLUTION_FILL = "EAF4EC"
LINK_BLUE = "0563C1"

STATUS_COLORS = {
    GLPI_STATUS[1]: "2F6DB5",  # new
    GLPI_STATUS[2]: "00838F",  # processing (assigned)
    GLPI_STATUS[3]: "00838F",  # processing (planned)
    GLPI_STATUS[4]: "B26A00",  # pending
    GLPI_STATUS[5]: "2E7D32",  # solved
    GLPI_STATUS[6]: "616161",  # closed
}
DEFAULT_STATUS_COLOR = "616161"

MONTHS_PT = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]

PARAGRAPH_SEPARATOR = "\n\n"  # models.html_to_text separates paragraphs with a blank line

# A text is only treated as junk if it is long AND has a giant "word" with
# no spaces (e.g. an image pasted as base64 without the usual prefix).
# Length alone is not enough — it once cut a legitimate 3,000-character solution.
TEXT_LENGTH_LIMIT = 1500
NO_SPACE_RUN_LIMIT = 200


def build_report(
    tickets,
    output_path: str,
    images_by_document: dict | None = None,
    instance_name: str | None = None,
    technician_by_ticket: dict | None = None,
    period: tuple | None = None,
):
    """Generates a single .docx with all the given tickets.

    `images_by_document`: {GLPI_document_id: local_path} with the images
    already downloaded from GLPI (see glpi_files.py). An image attachment
    with no entry here is shown as "sem imagem disponível".

    `technician_by_ticket`: optional {ticket_id: technician_full_name} (see
    the --technician filter in generate_report.py). When given, the report
    is grouped by technician (one heading per person, before their tickets)
    — `tickets` is expected to be already sorted in that same order.

    `instance_name`: name of the source GLPI, shown on the cover and header.

    `period`: optional (start_date, end_date), shown on the cover and in the
    page header."""
    images_by_document = images_by_document or {}
    technician_by_ticket = technician_by_ticket or {}
    period_text = _period_text(period)

    doc = Document()
    _setup_page(doc)
    _setup_styles(doc)
    _setup_header_footer(doc, instance_name, period_text)
    _cover(doc, tickets, instance_name, technician_by_ticket, images_by_document, period_text)
    _table_of_contents(doc, tickets, technician_by_ticket)

    previous_technician = object()  # sentinel: guarantees a heading before the first group
    for ticket in tickets:
        # the first ticket of each group does not break the page — the
        # technician heading right before it already did (see _technician_heading)
        page_break = True
        if technician_by_ticket:
            current = technician_by_ticket.get(ticket.id)
            if current != previous_technician:
                _technician_heading(doc, current)
                previous_technician = current
                page_break = False
        _ticket_section(doc, ticket, images_by_document, page_break)

    _set_document_language(doc)
    _set_properties(doc, instance_name, period_text)
    doc.save(output_path)


# ---------------------------------------------------------------------------
# Page, styles, header and footer
# ---------------------------------------------------------------------------

def _setup_page(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(0.9)
    section.different_first_page_header_footer = True  # the cover has no header/footer


def _style_font(style, size_pt, bold=None, color_hex=None):
    style.font.name = FONT
    style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    if color_hex:
        style.font.color.rgb = RGBColor.from_string(color_hex)
    rfonts = style.element.rPr.rFonts
    # Word's built-in headings use a theme font, which overrides the name set above
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attribute), None)
    rfonts.set(qn("w:eastAsia"), FONT)
    rfonts.set(qn("w:cs"), FONT)


def _setup_styles(doc: Document):
    normal = doc.styles["Normal"]
    _style_font(normal, 10.5, color_hex="262626")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.widow_control = True  # no single orphan/widow line at a page edge

    h1 = doc.styles["Heading 1"]  # technician name
    _style_font(h1, 18, bold=True, color_hex=ACCENT)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]  # ticket
    _style_font(h2, 14, bold=True, color_hex=ACCENT)
    h2.paragraph_format.space_before = Pt(0)
    h2.paragraph_format.space_after = Pt(8)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]  # Description / Updates / Attachments
    _style_font(h3, 11.5, bold=True, color_hex=ACCENT)
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True


def _setup_header_footer(doc: Document, instance_name: str | None, period_text: str):
    label = "Book de Evidências" + (f" — {instance_name}" if instance_name else "")
    section = doc.sections[0]

    # header (page 2 onward): report name on the left, period on the right
    p = section.header.paragraphs[0]
    p.style = doc.styles["Normal"]
    p.paragraph_format.tab_stops.add_tab_stop(Cm(USABLE_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    _run(p, label, 8.5, MUTED_TEXT)
    if period_text:
        _run(p, "\t" + period_text, 8.5, MUTED_TEXT)
    _paragraph_border(p, "bottom", SOFT_BORDER, sz=4, space=3)

    # footer: generation date and "Página X de Y", all in the same size
    p = section.footer.paragraphs[0]
    p.style = doc.styles["Normal"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, f"Gerado em {datetime.now().strftime('%d/%m/%Y')}   |   Página ", 8.5, MUTED_TEXT)
    _insert_field(p, "PAGE", 8.5, MUTED_TEXT)
    _run(p, " de ", 8.5, MUTED_TEXT)
    _insert_field(p, "NUMPAGES", 8.5, MUTED_TEXT)


# ---------------------------------------------------------------------------
# Formatting helpers (Word XML)
# ---------------------------------------------------------------------------

def _run(paragraph, text, size_pt=None, color_hex=None, bold=False, italic=False):
    r = paragraph.add_run(text)
    if size_pt:
        r.font.size = Pt(size_pt)
    if color_hex:
        r.font.color.rgb = RGBColor.from_string(color_hex)
    r.bold = bold or None
    r.italic = italic or None
    return r


def _badge(paragraph, text, color_hex, size_pt=9):
    """Bold white text on a colored background (a "pill")."""
    r = _run(paragraph, f" {text} ", size_pt, "FFFFFF", bold=True)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    r._r.get_or_add_rPr().insert_element_before(
        shd, "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath"
    )
    return r


# Elements that must come AFTER <w:pBdr> inside <w:pPr> (OOXML schema order)
_PBDR_SUCCESSORS = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
    "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def _paragraph_border(paragraph, side: str, color_hex: str, sz: int = 8, space: int = 4):
    """A rule on one side of the paragraph (`side`: top/left/bottom/right).
    `sz` in eighths of a point."""
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.insert_element_before(pbdr, *_PBDR_SUCCESSORS)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(sz))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color_hex)
    pbdr.append(border)


def _paragraph_fill(paragraph, color_hex: str):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    ppr.insert_element_before(shd, *_PBDR_SUCCESSORS[1:])


def _cell_fill(cell, color_hex: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcpr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


def _table_borders(table, color_hex=SOFT_BORDER, sz=4, horizontal=True, vertical=False, outer=True):
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name, enabled in (
        ("top", outer), ("left", outer and vertical), ("bottom", outer),
        ("right", outer and vertical), ("insideH", horizontal), ("insideV", vertical),
    ):
        b = OxmlElement(f"w:{name}")
        b.set(qn("w:val"), "single" if enabled else "nil")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), color_hex)
        borders.append(b)
    tblpr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")


def _column_widths(table, widths_cm: list):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, width in zip(table.columns, widths_cm):
        column.width = Cm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def _keep_row_together(row, repeat_as_header=False):
    """The row never splits across pages; optionally repeats as the table
    header at the top of every page."""
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))
    if repeat_as_header:
        trpr.append(OxmlElement("w:tblHeader"))


def _cell_text(cell, text="", size_pt=10, color_hex=None, bold=False, align=None, spacing_pt=2):
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(spacing_pt)
    p.paragraph_format.space_after = Pt(spacing_pt)
    if align is not None:
        p.alignment = align
    if text:
        _run(p, text, size_pt, color_hex, bold)
    return p


def _insert_field(paragraph, instruction: str, size_pt=10, color_hex=MUTED_TEXT):
    """A Word dynamic field (PAGE, NUMPAGES...). A field is a sequence of
    runs: begin -> instruction -> separate -> placeholder value -> end.
    Each run gets the same size/color — with a single run, the field
    result showed up in the default size, bigger than the text around it."""
    def _new_run():
        r = paragraph.add_run()
        r.font.size = Pt(size_pt)
        r.font.color.rgb = RGBColor.from_string(color_hex)
        return r

    r = _new_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    r._r.append(begin)

    r = _new_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    r._r.append(instr)

    r = _new_run()
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    r._r.append(separator)

    r = _new_run()
    r.text = "1"  # placeholder; Word/LibreOffice recompute it

    r = _new_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.append(end)


def _set_document_language(doc: Document, language: str = "pt-BR"):
    """Marks every run already in the document as pt-BR — otherwise Word
    assumes the operating system language (e.g. English) and underlines
    the whole text in red as misspelled. Must run last, after all content
    has been added (it only tags runs that already exist)."""
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
        paragraphs.extend(section.first_page_footer.paragraphs)
        paragraphs.extend(section.first_page_header.paragraphs)

    for paragraph in paragraphs:
        # recursive search (not just paragraph.runs) because the runs of the
        # internal links live inside <w:hyperlink>, not directly in the paragraph
        for r_el in paragraph._p.findall(".//" + qn("w:r")):
            rpr = r_el.get_or_add_rPr()
            lang = OxmlElement("w:lang")
            lang.set(qn("w:val"), language)
            lang.set(qn("w:eastAsia"), language)
            lang.set(qn("w:bidi"), language)
            rpr.append(lang)


def _set_properties(doc: Document, instance_name: str | None, period_text: str):
    """File properties (File > Info in Word, and the PDF metadata).
    Without this, the .docx carries the python-docx template's author."""
    props = doc.core_properties
    props.title = "Book de Evidências" + (f" — {instance_name}" if instance_name else "")
    props.subject = period_text or ""
    props.author = "glpi-evidence-report"
    props.last_modified_by = "glpi-evidence-report"
    props.comments = ""
    props.keywords = "GLPI, evidências, atendimentos"
    now = datetime.now()
    props.created = props.modified = now
    props.revision = 1


def _sanitize_text(text: str) -> str:
    """Keeps a giant chunk of junk data (e.g. someone pasted an image
    straight into a GLPI description/followup and it was stored as raw
    base64 text instead of a real attachment) from turning into pages of
    meaningless text in the report.

    Base64 image: returns empty (it is not real evidence, no trace is
    left). Long chunk with a huge run without spaces: replaced by a short
    note, since there may be something relevant there, just badly
    formatted. Normal long text passes through untouched."""
    if not text:
        return text
    looks_like_base64_image = "base64," in text.lower() or "data:image" in text.lower()
    if looks_like_base64_image:
        return ""
    if len(text) > TEXT_LENGTH_LIMIT:
        longest_run = max((len(t) for t in text.split()), default=0)
        if longest_run > NO_SPACE_RUN_LIMIT:
            return f"[Trecho muito longo omitido ({len(text)} caracteres) — provavelmente dado inválido no campo]"
    return text


def _ticket_anchor(ticket_id: int) -> str:
    return f"chamado_{ticket_id}"


def _add_bookmark(paragraph, name: str, id_num: int):
    """Marks a paragraph with a bookmark (internal reference) — used on each
    ticket title, so the table of contents can link straight to it."""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(id_num))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(id_num))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _add_internal_link(paragraph, text: str, anchor: str, size_pt: float = 10):
    """Internal hyperlink (jumps to a bookmark in the same document) —
    unlike Word's native TOC field, it works the same in Word desktop, Word
    for the web and in the PDF, with no "Update field" step."""
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_BLUE)
    rpr.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(int(size_pt * 2)))  # w:sz is in half-points
    rpr.append(size)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)

    paragraph._p.append(hyperlink)


# ---------------------------------------------------------------------------
# Cover and table of contents
# ---------------------------------------------------------------------------

def _period_text(period) -> str:
    if not period:
        return ""
    start, end = period
    if isinstance(start, date) and isinstance(end, date):
        whole_month = (
            start.day == 1
            and start.year == end.year
            and start.month == end.month
            and end.day == calendar.monthrange(end.year, end.month)[1]
        )
        if whole_month:
            return f"{MONTHS_PT[start.month - 1]} de {start.year}"
        return f"{start:%d/%m/%Y} a {end:%d/%m/%Y}"
    return ""


def _count_evidence(ticket, images_by_document: dict) -> int:
    return sum(1 for a in ticket.attachments if a.is_image and a.id in images_by_document)


def _cover(doc, tickets, instance_name, technician_by_ticket, images_by_document, period_text):
    # space at the top, so the cover does not stick to the margin
    doc.add_paragraph().paragraph_format.space_after = Pt(60)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    _run(p, "BOOK DE EVIDÊNCIAS", 30, ACCENT, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    _run(p, "Atendimentos registrados no GLPI", 16, MUTED_TEXT)
    _paragraph_border(p, "bottom", ACCENT, sz=18, space=8)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    if instance_name:
        _run(p, f"Sistema: {instance_name}", 13, "262626", bold=True)
    if period_text:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        _run(p, f"Período: {period_text}", 13, "262626", bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(26)
    _run(p, f"Gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M')}", 10.5, MUTED_TEXT)

    # summary per technician (or a single total, if the report is not grouped)
    groups = {}
    for t in tickets:
        name = technician_by_ticket.get(t.id) if technician_by_ticket else None
        groups.setdefault(name or "Total", []).append(t)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    _run(p, "Resumo", 12, ACCENT, bold=True)

    columns = ["Técnico" if technician_by_ticket else "", "Chamados", "Evidências (imagens)", "Solucionados / fechados", "Em aberto"]
    table = doc.add_table(rows=1, cols=len(columns))
    for i, (cell, text) in enumerate(zip(table.rows[0].cells, columns)):
        _cell_fill(cell, ACCENT)
        _cell_text(cell, text, 9.5, "FFFFFF", True, WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER, 4)
    _keep_row_together(table.rows[0], repeat_as_header=True)

    def _row(name, group, highlight=False):
        resolved = sum(1 for t in group if t.status in (STATUS_SOLVED, STATUS_CLOSED))
        values = [
            name, str(len(group)), str(sum(_count_evidence(t, images_by_document) for t in group)),
            str(resolved), str(len(group) - resolved),
        ]
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, values)):
            if highlight:
                _cell_fill(cell, LABEL_FILL)
            _cell_text(cell, value, 10, None, highlight, WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER, 4)

    for name in sorted(groups):
        _row(name, groups[name])
    if len(groups) > 1:
        _row("Total", list(tickets), highlight=True)
    _table_borders(table)
    _column_widths(table, [6.2, 2.4, 3.2, 3.2, 2.0])


def _table_of_contents(doc, tickets, technician_by_ticket):
    title = doc.add_heading("Sumário dos atendimentos", level=1)
    title.paragraph_format.page_break_before = True

    widths = [1.7, 10.3, 2.9, 2.1]
    table = doc.add_table(rows=1, cols=4)
    for cell, text in zip(table.rows[0].cells, ["Chamado", "Título", "Situação", "Abertura"]):
        _cell_fill(cell, ACCENT)
        _cell_text(cell, text, 9, "FFFFFF", True, spacing_pt=3)
    _keep_row_together(table.rows[0], repeat_as_header=True)

    previous_technician = object()
    for t in tickets:
        if technician_by_ticket:
            current = technician_by_ticket.get(t.id)
            if current != previous_technician:
                row = table.add_row()
                cell = row.cells[0].merge(row.cells[3])
                _cell_fill(cell, LABEL_FILL)
                _cell_text(cell, current or "Outros atendimentos", 10, ACCENT, True, spacing_pt=4)
                _keep_row_together(row)
                cell.paragraphs[0].paragraph_format.keep_with_next = True
                previous_technician = current

        row = table.add_row()
        _keep_row_together(row)
        c0, c1, c2, c3 = row.cells
        _add_internal_link(_cell_text(c0), f"#{t.id}", _ticket_anchor(t.id), 9)
        _cell_text(c1, t.title, 9)
        _cell_text(c2, t.status.replace("Em atendimento", "Em atend."), 9, STATUS_COLORS.get(t.status, DEFAULT_STATUS_COLOR), True)
        _cell_text(c3, _fmt_day(t.opened_at), 9)
    _table_borders(table)
    _column_widths(table, widths)


def _technician_heading(doc: Document, technician_name: str | None):
    """Section divider before each technician's tickets, when the report is
    grouped by --technician (see generate_report.py). The divider itself
    breaks the page, and the group's first ticket comes without a break
    (`page_break=False` in _ticket_section) — so both stay together at the
    top of the new page. If the break were only on the ticket, the
    technician's name would be left alone at the bottom of the previous page."""
    heading = doc.add_heading(technician_name or "Outros atendimentos", level=1)
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.space_after = Pt(14)


# ---------------------------------------------------------------------------
# One ticket
# ---------------------------------------------------------------------------

def _fmt_datetime(value) -> str:
    if not value:
        return "-"
    try:
        return value.strftime("%d/%m/%Y %H:%M")
    except AttributeError:
        return str(value)


def _fmt_day(value) -> str:
    if not value:
        return "-"
    try:
        return value.strftime("%d/%m/%Y")
    except AttributeError:
        return str(value)


def _ticket_section(doc: Document, ticket, images_by_document: dict, page_break: bool = True):
    title = doc.add_heading(f"#{ticket.id} — {ticket.title}", level=2)
    title.paragraph_format.page_break_before = page_break
    _paragraph_border(title, "bottom", ACCENT, sz=12, space=3)
    _add_bookmark(title, _ticket_anchor(ticket.id), ticket.id)

    _metadata_grid(doc, ticket)

    doc.add_heading("Descrição", level=3)
    for paragraph in (ticket.description or "-").split(PARAGRAPH_SEPARATOR):
        clean = _sanitize_text(paragraph)
        if clean:
            doc.add_paragraph(clean)

    if ticket.updates:
        doc.add_heading("Atualizações / Acompanhamentos", level=3)
        for update in ticket.updates:
            _update(doc, update)

    if ticket.attachments:
        doc.add_heading("Anexos", level=3)
        _insert_attachments(doc, ticket.attachments, images_by_document)


def _metadata_grid(doc: Document, ticket):
    """4 x 2 grid (label | value | label | value) in shades of gray, with the
    status as a colored badge."""
    pairs = [
        (("Situação", None), ("Prioridade", ticket.priority)),
        (("Categoria", ticket.category), ("Entidade", ticket.entity)),
        (("Solicitante(s)", ", ".join(ticket.requesters()) or "-"), ("Responsável(is)", ", ".join(ticket.assignees()) or "-")),
        (("Abertura", _fmt_datetime(ticket.opened_at)), ("Solução", _fmt_datetime(ticket.solved_at))),
    ]
    table = doc.add_table(rows=0, cols=4)
    for (label_a, value_a), (label_b, value_b) in pairs:
        cells = table.add_row().cells
        _keep_row_together(table.rows[-1])
        for label_cell, label in ((cells[0], label_a), (cells[2], label_b)):
            _cell_fill(label_cell, LABEL_FILL)
            _cell_text(label_cell, label, 8.5, MUTED_TEXT, True, spacing_pt=3)
        _cell_text(cells[1], value_a if value_a is not None else "", 10, spacing_pt=3)
        _cell_text(cells[3], value_b, 10, spacing_pt=3)
        if label_a == "Situação":
            _badge(cells[1].paragraphs[0], ticket.status, STATUS_COLORS.get(ticket.status, DEFAULT_STATUS_COLOR))
    _table_borders(table, horizontal=True, vertical=False)
    _column_widths(table, [2.6, 5.9, 2.6, 5.9])


def _update(doc: Document, update):
    """One update as a timeline entry: date and author on a small line, the
    text right below. The solution gets a badge, a side bar and a green
    background."""
    info = doc.add_paragraph()
    info.paragraph_format.space_before = Pt(6)
    info.paragraph_format.space_after = Pt(1)
    info.paragraph_format.keep_with_next = True
    if update.is_solution:
        _badge(info, "SOLUÇÃO", SOLUTION_GREEN, 8)
        _run(info, "  ", 9)
    _run(info, f"{_fmt_datetime(update.date)}  ·  {update.author}", 9, MUTED_TEXT, bold=True)

    body = doc.add_paragraph(_sanitize_text(update.text) or "-")
    body.paragraph_format.space_after = Pt(4)
    if update.is_solution:
        for p in (info, body):
            p.paragraph_format.left_indent = Cm(0.3)
            _paragraph_border(p, "left", SOLUTION_GREEN, sz=24, space=6)
            _paragraph_fill(p, SOLUTION_FILL)


# ---------------------------------------------------------------------------
# Attachments
# ---------------------------------------------------------------------------

def _origin_caption(attachment) -> str:
    """Caption text saying where the image came from, e.g.
    "Solução de 12/08/2026, Ana Ribeiro"."""
    parts = [f"{attachment.origin} de {_fmt_day(attachment.origin_date)}".strip()]
    if attachment.origin_author:
        parts.append(attachment.origin_author)
    return ", ".join(parts)


def _insert_attachments(doc: Document, attachments, images_by_document: dict):
    non_images = []  # labels of attachments that are not images
    resolved_images = []  # [(attachment, path), ...]
    missing = []  # names of image attachments with no downloaded file

    for attachment in attachments:
        if not attachment.is_image:
            label = attachment.name
            if attachment.filename and attachment.filename != attachment.name:
                label = f"{attachment.name} — {attachment.filename}"
            non_images.append(label)
            continue

        path = images_by_document.get(attachment.id)
        if path is None:
            missing.append(attachment.name)
            continue
        resolved_images.append((attachment, path))

    for label in non_images:
        doc.add_paragraph(f"📎 {label}")

    total = len(resolved_images)
    for i, (attachment, path) in enumerate(resolved_images, start=1):
        try:
            _insert_scaled_image(doc, path)
        except Exception as exc:  # noqa: BLE001 — one broken image must not stop the whole report
            doc.add_paragraph(f"[Erro ao inserir imagem: {exc}]")
            continue
        image_paragraph = doc.paragraphs[-1]
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_before = Pt(6)
        image_paragraph.paragraph_format.space_after = Pt(2)
        # "keep with next": the image stays on the same page as its caption
        # (otherwise the caption can end up alone on the next page)
        image_paragraph.paragraph_format.keep_with_next = True

        # caption below the image, as in technical report figures
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(12)
        _run(caption, f"Evidência {i} de {total}", 9, "262626", bold=True)
        origin = _origin_caption(attachment)
        if origin:
            _run(caption, f" — {origin}", 9, MUTED_TEXT)

    for name in missing:
        doc.add_paragraph(f'[Anexo "{name}" sem imagem disponível]')


def _insert_scaled_image(doc: Document, path: str):
    """Inserts the image bounded by the most restrictive dimension (width OR
    height), not just width — keeps a very tall/narrow screenshot (e.g. a
    whole page captured vertically) from scaling proportionally and taking
    several pages on its own. The image itself is not changed; it only
    gets a thin frame."""
    try:
        image = DocxImage.from_file(path)
        height_at_max_width = MAX_IMAGE_WIDTH_CM * (image.height / image.width)
    except Exception:  # noqa: BLE001 — could not read dimensions (unusual file): fall back to the default
        doc.add_picture(path, width=Cm(MAX_IMAGE_WIDTH_CM))
        _frame_image(doc)
        return

    if height_at_max_width > MAX_IMAGE_HEIGHT_CM:
        doc.add_picture(path, height=Cm(MAX_IMAGE_HEIGHT_CM))
    else:
        doc.add_picture(path, width=Cm(MAX_IMAGE_WIDTH_CM))
    _frame_image(doc)


def _frame_image(doc: Document):
    """Thin gray outline on the last inserted image (DrawingML, `a:` prefix)."""
    sp_pr = doc.inline_shapes[-1]._inline.graphic.graphicData.pic.spPr
    outline = parse_xml(
        f'<a:ln {nsdecls("a")} w="9525"><a:solidFill><a:srgbClr val="BFBFBF"/></a:solidFill></a:ln>'
    )
    sp_pr.append(outline)
